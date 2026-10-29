#!/usr/bin/env python3
"""
sync_leetcode.py - 把你在 LeetCode 上已经 Accepted 的题,自动在 Excel 刷题计划里划掉并更新进度。

用法(三种来源,可叠加):
  1) 粘贴你的 LeetCode 做题记录(复制 Progress / Submissions 页面的文字存成 solved.txt):
       python sync_leetcode.py --xlsx NeetCode150_刷题计划.xlsx --paste solved.txt
  2) 登录态(一次拿到全部已做题,推荐):
       export LEETCODE_SESSION="浏览器 Cookie 里的 LEETCODE_SESSION 值"
       python sync_leetcode.py --xlsx NeetCode150_刷题计划.xlsx --cookie-env LEETCODE_SESSION
  3) 公开主页(无需登录,但 LeetCode 只返回最近 20 条 Accepted,适合每天跑一次累积):
       python sync_leetcode.py --xlsx NeetCode150_刷题计划.xlsx --username 你的用户名

常用选项:  --dry-run 只显示会改什么,不写文件;  --next 3 显示接下来要做的 3 题。
每次运行会:备份原表为 *.bak.xlsx → 划掉已完成题(删除线+灰字+绿底)→ 补全"完成日期"→ 刷新进度。
(NeetCode 模板里"完成"原本只认用时列;脚本会把进度公式改成"填了一刷日期或用时都算完成"。)
重复运行是安全的(已划掉的不会再改)。已做题记录累积保存在 solved_state.json。

依赖: Python 3.9+, openpyxl (pip install openpyxl)。联网只用标准库。
"""
import argparse, json, os, re, shutil, sys, urllib.request
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

DONE_FILL = PatternFill("solid", fgColor="E2F0D9")
GRAY = "7F7F7F"
UA = {"User-Agent": "Mozilla/5.0 (leetcode-sync)", "Referer": "https://leetcode.com/"}

# ---------------------------------------------------------------- sources
LABEL = re.compile(r"^(Today|Yesterday|Mon|Tue|Wed|Thu|Fri|Sat|Sun|[A-Z][a-z]{2}\s+\d{1,2})$")
LINK = re.compile(r"\[\s*(\d{1,4})\.\s*([^\]]+?)\s*\]\(\s*https?://leetcode\.(?:com|cn)/problems/([a-z0-9\-]+)")
PLAIN = re.compile(r"^\s*(\d{1,4})\.\s+([A-Za-z].*?)\s*$")
WD = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

def resolve_label(label, today):
    """LeetCode's relative labels ('Yesterday', 'Mon', 'Oct 19') -> ISO date string, or None."""
    if not label: return None
    if label == "Today": return today.isoformat()
    if label == "Yesterday": return (today - timedelta(days=1)).isoformat()
    if label in WD:
        delta = (today.weekday() - WD.index(label)) % 7 or 7
        return (today - timedelta(days=delta)).isoformat()
    m = re.match(r"([A-Z][a-z]{2})\s+(\d{1,2})$", label)
    if m and m.group(1) in MON:
        d = date(today.year, MON.index(m.group(1)) + 1, int(m.group(2)))
        if d > today: d = d.replace(year=today.year - 1)
        return d.isoformat()
    return None

def parse_paste(text, today=None):
    """Return {number: {"title","slug","date"}} from pasted LeetCode progress text.
    Accepts "[1068. Title](https://leetcode.com/problems/slug/)" links and plain "1068. Title" lines;
    the relative date label above each item ('Yesterday', 'Mon', 'Oct 19') is converted to a real date."""
    today = today or date.today()
    out, cur = {}, None
    for line in text.splitlines():
        t = line.strip()
        if LABEL.match(t):
            cur = resolve_label(t, today); continue
        m = LINK.search(t)
        if m:
            out.setdefault(int(m.group(1)), {"title": m.group(2), "slug": m.group(3), "date": cur}); continue
        m = PLAIN.match(t)
        if m:
            out.setdefault(int(m.group(1)), {"title": m.group(2), "slug": None, "date": cur})
    return out

def _get(url, cookie=None, data=None, headers=None):
    h = dict(UA)
    if headers: h.update(headers)
    if cookie: h["Cookie"] = cookie
    req = urllib.request.Request(url, data=data, headers=h)
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))

def fetch_problem_index(session=None):
    """slug -> (frontend number, title, status). status is only meaningful when logged in."""
    cookie = f"LEETCODE_SESSION={session}" if session else None
    data = _get("https://leetcode.com/api/problems/all/", cookie=cookie)
    idx = {}
    for p in data.get("stat_status_pairs", []):
        st = p["stat"]
        idx[st["question__title_slug"]] = (st["frontend_question_id"], st["question__title"], p.get("status"))
    return idx

def fetch_logged_in(session):
    idx = fetch_problem_index(session)
    return {n: {"title": t, "slug": s} for s, (n, t, st) in idx.items() if st == "ac"}

def fetch_public(username, limit=20):
    q = {"query": "query r($u:String!,$l:Int!){recentAcSubmissionList(username:$u,limit:$l){title titleSlug timestamp}}",
         "variables": {"u": username, "l": limit}}
    res = _get("https://leetcode.com/graphql", data=json.dumps(q).encode(), headers={"Content-Type": "application/json"})
    subs = (res.get("data") or {}).get("recentAcSubmissionList") or []
    idx = fetch_problem_index()
    out = {}
    for s in subs:
        if s["titleSlug"] in idx:
            n, t, _ = idx[s["titleSlug"]]
            d = datetime.fromtimestamp(int(s["timestamp"])).date().isoformat()
            out.setdefault(n, {"title": t, "slug": s["titleSlug"], "date": d})
    return out

# ---------------------------------------------------------------- state
def load_state(path):
    if path.exists():
        return {int(k): v for k, v in json.loads(path.read_text(encoding="utf-8")).items()}
    return {}

def save_state(path, state):
    path.write_text(json.dumps({str(k): v for k, v in sorted(state.items())}, ensure_ascii=False, indent=1), encoding="utf-8")

# ---------------------------------------------------------------- excel
NUM_HDR = re.compile(r"题号|题目编号|problem\s*(no|num|#|id)|^#$|^no\.?$|^id$|number", re.I)
DATE_HDR = re.compile(r"完成日期|做题日期|日期|date", re.I)
STATUS_HDR = re.compile(r"^(完成|状态|已完成|status|done)", re.I)
SKIP_HDR = re.compile(r"二刷|second|计划|plan|due", re.I)

def lead_int(v):
    if v is None: return None
    if isinstance(v, (int, float)) and 1 <= int(v) <= 4000: return int(v)
    m = re.match(r"^\s*#?(\d{1,4})\b", str(v))
    return int(m.group(1)) if m else None

def find_layout(ws, allow_fallback=True):
    """Find header row + number column + optional date/status columns."""
    for r in range(1, min(ws.max_row, 12) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and NUM_HDR.search(v.strip()):
                lay = {"hdr": r, "num": c, "date": None, "status": None, "title": None, "diff": None, "note": None}
                for cc in range(1, ws.max_column + 1):
                    h = ws.cell(r, cc).value
                    if not isinstance(h, str): continue
                    h = h.strip()
                    if lay["note"] is None and re.search(r"备注|note", h, re.I): lay["note"] = cc
                    if lay["diff"] is None and re.search(r"难度|difficulty", h, re.I): lay["diff"] = cc
                    if lay["title"] is None and re.search(r"题名|题目名|title|name", h, re.I) and cc != c: lay["title"] = cc
                    if SKIP_HDR.search(h) and not re.search(r"完成日期", h): continue
                    if lay["date"] is None and DATE_HDR.search(h): lay["date"] = cc
                    elif lay["status"] is None and STATUS_HDR.search(h): lay["status"] = cc
                return lay
    # fallback: column with most integer-like cells
    best, bc = 0, None
    for c in range(1, ws.max_column + 1):
        k = sum(1 for r in range(1, ws.max_row + 1) if lead_int(ws.cell(r, c).value))
        if k > best: best, bc = k, c
    if allow_fallback and bc and best >= 5:
        return {"hdr": 0, "num": bc, "date": None, "status": None, "title": None, "diff": None, "note": None}
    return None

def is_struck(cell):
    return bool(cell.font and cell.font.strike)

def strike_row(ws, r, maxc):
    for c in range(1, maxc + 1):
        cell = ws.cell(r, c)
        f = copy(cell.font)
        f.strike = True
        f.color = Color(rgb="FF" + GRAY)
        cell.font = f
        cell.fill = DONE_FILL

def as_cell_date(d):
    try: return date.fromisoformat(d)
    except Exception: return d

def sumdone(end):
    return f'(((刷题记录!$F$2:$F${end}<>"")+(刷题记录!$G$2:$G${end}<>""))>0)'

def patch_progress(wb):
    """Template counts a problem as done only when 用时(G) is filled. Make the 'done' counts also accept 一刷日期(F).
    Works for any plan length (reads the end row from the formula)."""
    if "刷题记录" not in wb.sheetnames: return False
    n = 0
    for name in ("进度总览", "每周计划"):
        if name not in wb.sheetnames: continue
        for row in wb[name].iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=") and "SUMPRODUCT" not in v): continue
                m = re.fullmatch(r"=COUNT\(刷题记录!\$G\$2:\$G\$(\d+)\)", v)
                if m:
                    c.value = "=SUMPRODUCT(--" + sumdone(m.group(1)) + ")"; n += 1; continue
                m = re.fullmatch(r'=COUNTIFS\((刷题记录!\$[A-Z]\$2:\$[A-Z]\$(\d+)),([A-Z]+\d+),刷题记录!\$G\$2:\$G\$\d+,">0"\)', v)
                if m:
                    c.value = f"=SUMPRODUCT(({m.group(1)}={m.group(3)})*" + sumdone(m.group(2)) + ")"; n += 1
    return n > 0

def update_workbook(xlsx, solved, today, dry_run=False, next_n=3):
    wb = load_workbook(xlsx)
    plan_nums, newly, already, pending = set(), [], [], []
    layouts = {ws.title: find_layout(ws, allow_fallback=False) for ws in wb.worksheets if ws.title not in ("进度", "进度总览")}
    if not any(layouts.values()):
        layouts = {ws.title: find_layout(ws) for ws in wb.worksheets if ws.title != "进度"}
    for ws in wb.worksheets:
        lay = layouts.get(ws.title)
        if not lay: continue
        for r in range(lay["hdr"] + 1, ws.max_row + 1):
            n = lead_int(ws.cell(r, lay["num"]).value)
            if n is None: continue
            plan_nums.add(n)
            title = (str(ws.cell(r, lay["title"]).value) if lay.get("title") else
                     next((str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)
                           if c != lay["num"] and isinstance(ws.cell(r, c).value, str) and len(str(ws.cell(r, c).value)) > 3), ""))
            if n in solved:
                if is_struck(ws.cell(r, lay["num"])):
                    already.append(n)
                else:
                    newly.append((n, ws.title, r))
                    if not dry_run:
                        strike_row(ws, r, ws.max_column)
                        d = as_cell_date(solved[n].get("date") or today)
                        if lay["date"] and ws.cell(r, lay["date"]).value in (None, ""):
                            ws.cell(r, lay["date"]).value = d
                        if lay["status"]:
                            ws.cell(r, lay["status"]).value = "✓"
            else:
                pending.append((n, title, (str(ws.cell(r, lay["diff"]).value or "") if lay.get("diff") else "")
                               + ("会员题" if lay.get("note") and "会员" in str(ws.cell(r, lay["note"]).value or "") else "")))
    total, done = len(plan_nums), len(plan_nums & set(solved))
    extra = sorted(set(solved) - plan_nums)
    if not dry_run:
        if patch_progress(wb) or "进度总览" in wb.sheetnames:
            wb.save(xlsx)
            return {"total": total, "done": done, "newly": newly, "already": already,
                    "extra": extra, "next": pending[:next_n]}
        if "进度" in wb.sheetnames: del wb["进度"]
        ps = wb.create_sheet("进度", 0)
        rows = [("总题数(计划)", total), ("已完成", done), ("剩余", total - done),
                ("完成率", f"{done / total:.0%}" if total else "-"), ("更新日期", today), ("", ""),
                ("已做但不在计划里", ", ".join(map(str, extra)) or "无")]
        for i, (a, b) in enumerate(rows, 1):
            ps.cell(i, 1).value, ps.cell(i, 2).value = a, b
        ps.column_dimensions["A"].width = 20; ps.column_dimensions["B"].width = 60
        wb.save(xlsx)
    return {"total": total, "done": done, "newly": newly, "already": already,
            "extra": extra, "next": pending[:next_n]}

# ---------------------------------------------------------------- reminder text
def slug_of(title):
    return re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")

def build_summary(res, state, today, mode, warnings):
    pct = f"{res['done']}/{res['total']}({res['done'] / res['total']:.0%})" if res["total"] else "-"
    done_today = sorted(n for n, v in state.items() if v.get("date") == today)
    lines = []
    if mode == "night":
        if done_today:
            lines.append(f"✅ 今天做完 {len(done_today)} 题:{', '.join(map(str, done_today))}")
        else:
            lines.append("⏰ 今天还没做题,睡前至少做 1 题?")
    else:
        lines.append("📚 今日刷题(目标 3 题)")
    if mode == "morning" or not done_today:
        for n, t, d in res["next"]:
            tag = f" [{d.replace('会员题', '/会员题,可去 neetcode.io 做') if d.endswith('会员题') and d != '会员题' else d}]" if d else ""
            lines.append(f"• {n}. {t}{tag}\n  https://leetcode.com/problems/{slug_of(t)}/")
    lines.append(f"进度 {pct}")
    lines += [f"⚠ {w}" for w in warnings]
    return "\n".join(lines)

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Cross off solved LeetCode problems in an Excel plan.")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--paste", help="text file with pasted LeetCode progress")
    ap.add_argument("--cookie-env", help="env var holding LEETCODE_SESSION cookie")
    ap.add_argument("--username", help="public LeetCode username (last 20 AC only)")
    ap.add_argument("--state", help="state json path (default: next to xlsx)")
    ap.add_argument("--pre-plan", help="comma-separated problem numbers done before the plan (no date), e.g. 217,1,49")
    ap.add_argument("--summary-file", help="write a reminder/recap message to this file")
    ap.add_argument("--mode", choices=["morning", "night"], default="morning", help="message style for --summary-file")
    ap.add_argument("--next", type=int, default=3)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    xlsx = Path(a.xlsx)
    if not xlsx.exists(): sys.exit(f"找不到 {xlsx}")
    state_path = Path(a.state) if a.state else xlsx.with_name("solved_state.json")
    state = load_state(state_path)
    today = date.today().isoformat()
    fetched, warnings = {}, []

    def attempt(label, fn):
        try:
            got = fn()
            if not got: warnings.append(f"{label} 没返回任何已做题")
            fetched.update(got or {})
            return bool(got)
        except Exception as e:  # network / auth / format problems: keep going with saved state
            warnings.append(f"{label} 同步失败:{type(e).__name__}: {e}")
            return False

    if a.pre_plan:
        for x in a.pre_plan.split(","):
            if x.strip(): fetched[int(x)] = {"title": None, "slug": None, "date": "计划前已做"}
    if a.paste: attempt("粘贴记录", lambda: parse_paste(Path(a.paste).read_text(encoding="utf-8")))
    ok = False
    if a.cookie_env:
        tok = os.environ.get(a.cookie_env)
        if tok: ok = attempt("登录态(LEETCODE_SESSION 可能已过期,请更新 Secret)", lambda: fetch_logged_in(tok))
        else: warnings.append(f"环境变量 {a.cookie_env} 为空")
    if a.username and not ok:
        attempt("公开主页", lambda: fetch_public(a.username))
    for n, v in fetched.items():
        if n not in state:
            state[n] = {"title": v.get("title"), "slug": v.get("slug"), "date": v.get("date") or today}
    if not state:
        sys.exit("没有任何已做题数据。请用 --paste / --cookie-env / --username 之一。\n" + "\n".join(warnings))
    if not a.dry_run:
        shutil.copy(xlsx, xlsx.with_suffix(".bak.xlsx"))
    res = update_workbook(xlsx, state, today, a.dry_run, a.next)
    if not a.dry_run: save_state(state_path, state)

    print(f"{'[预演] ' if a.dry_run else ''}进度:{res['done']}/{res['total']}({res['done']/res['total']:.0%})" if res["total"] else "表里没找到题号列")
    print("本次新划掉:", ", ".join(str(n) for n, _, _ in res["newly"]) or "无")
    if res["extra"]: print("已做但不在计划里:", ", ".join(map(str, res["extra"])))
    if res["next"]:
        print("接下来做:")
        for n, t, _ in res["next"]: print(f"  {n}. {t}")
    for w in warnings: print("⚠", w, file=sys.stderr)
    if a.summary_file:
        Path(a.summary_file).write_text(build_summary(res, state, today, a.mode, warnings), encoding="utf-8")

if __name__ == "__main__":
    main()
